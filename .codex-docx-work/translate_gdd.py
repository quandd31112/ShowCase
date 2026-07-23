from pathlib import Path

from docx import Document


SOURCE = Path("/Users/quan/Downloads/GDD-PRESSING A BUTTON.docx")
OUTPUT = Path("/Users/quan/Downloads/GDD-PRESSING A BUTTON-English.docx")

REPLACEMENTS = {
    "Người thực hiện": "Prepared by",
    (
        "Pressing a Button là game dạng survival funny 2D về một người tỉnh dậy mất hết ký ức khi bị đưa vào đây, "
        "trong 1 căn phòng trống với bức tường trắng với 1 cái nút lớn ở giữa phòng. Nhiệm vụ của player là ấn nút "
        "và sinh tồn trước những thứ mà cái nút đó đưa vào phòng. Player sẽ có 3 mạng"
    ): (
        "Pressing a Button is a humorous 2D survival game about an amnesiac who wakes up in an empty white room "
        "with a large button at its center. The player's objective is to press the button and survive whatever it "
        "brings into the room. The player has three lives."
    ),
    "Di chuyển": "Move",
    "Ấn nút": "Press the Button",
    "Sinh tồn trước những thứ mà khi ấn nút đưa đến": "Survive the hazards triggered by the button",
    "Input hiện tại": "Current Input",
    "Tương tác": "Interact",
    "Nhấn E": "Press E",
    "Nhấn Space": "Press Space",
    "2.3 Hệ thống Thành tựu": "2.3 Achievement System",
    (
        "Mỗi khi mở khóa thứ mới game sẽ thông báo cho người chơi và player có thể xem được tổng mục thành tựu"
    ): (
        "Whenever a new item or feature is unlocked, the game notifies the player. Players can review their complete "
        "achievement list at any time."
    ),
}


def replace_text(paragraph):
    original = paragraph.text
    replacement = REPLACEMENTS.get(original)
    if replacement is None:
        return False

    if paragraph.runs:
        paragraph.runs[0].text = replacement
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(replacement)
    return True


def replace_in_table(table):
    replacements = 0
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                replacements += replace_text(paragraph)
            for nested_table in cell.tables:
                replacements += replace_in_table(nested_table)
    return replacements


document = Document(SOURCE)
replacement_count = 0

for paragraph in document.paragraphs:
    replacement_count += replace_text(paragraph)

for table in document.tables:
    replacement_count += replace_in_table(table)

for section in document.sections:
    for container in (section.header, section.footer):
        for paragraph in container.paragraphs:
            replacement_count += replace_text(paragraph)
        for table in container.tables:
            replacement_count += replace_in_table(table)

document.save(OUTPUT)
print(f"Saved {OUTPUT} with {replacement_count} replacements")
