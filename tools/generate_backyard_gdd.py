from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path("/Users/quan/Documents/GitHub/ShowCase")
TEMPLATE = Path("/Users/quan/Downloads/CHICKCHICK-GDD.docx")
OUTPUT = ROOT / "KHU-VUON-MUA-HE-GDD.docx"


def clear_document_body(doc: Document) -> None:
    body = doc._body._element
    sect_pr = body.sectPr
    for child in list(body):
        if child is not sect_pr:
            body.remove(child)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_table_geometry(table, widths):
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")

    tbl_layout = tbl_pr.find(qn("w:tblLayout"))
    if tbl_layout is None:
        tbl_layout = OxmlElement("w:tblLayout")
        tbl_pr.append(tbl_layout)
    tbl_layout.set(qn("w:type"), "fixed")

    grid = tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            cell.width = Inches(widths[idx] / 1440)
            tc_w = cell._tc.get_or_add_tcPr().find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                cell._tc.get_or_add_tcPr().append(tc_w)
            tc_w.set(qn("w:w"), str(widths[idx]))
            tc_w.set(qn("w:type"), "dxa")
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cell)


def set_table_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        node = borders.find(qn(tag))
        if node is None:
            node = OxmlElement(tag)
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "4")
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), "D9DEE7")


def mark_header_row(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = tr_pr.find(qn("w:tblHeader"))
    if tbl_header is None:
        tbl_header = OxmlElement("w:tblHeader")
        tr_pr.append(tbl_header)
    tbl_header.set(qn("w:val"), "true")


def style_run(run, size=None, bold=None, color=None, font="Calibri"):
    run.font.name = font
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def add_para(doc, text="", style="normal", bold_prefix=None):
    p = doc.add_paragraph(style=style)
    if bold_prefix and text.startswith(bold_prefix):
        r1 = p.add_run(bold_prefix)
        style_run(r1, bold=True)
        r2 = p.add_run(text[len(bold_prefix):])
        style_run(r2)
    else:
        r = p.add_run(text)
        style_run(r)
    return p


def add_heading(doc, text, level, page_break_before=False):
    p = doc.add_heading(text, level=level)
    p.paragraph_format.page_break_before = page_break_before
    for run in p.runs:
        style_run(run, font="Calibri")
    return p


def add_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    mark_header_row(table.rows[0])
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = h
        set_cell_shading(hdr[i], "F2F4F7")
        for p in hdr[i].paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for r in p.runs:
                style_run(r, bold=True, color="0B2545")
    for row_data in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row_data):
            cells[i].text = value
            for p in cells[i].paragraphs:
                for r in p.runs:
                    style_run(r, size=10.5)
    set_table_borders(table)
    set_table_geometry(table, widths)
    doc.add_paragraph()
    return table


def set_document_defaults(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    for style_name in ("normal", "Normal"):
        if style_name in doc.styles:
            style = doc.styles[style_name]
            style.font.name = "Calibri"
            style._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
            style.font.size = Pt(11)
            style.paragraph_format.space_after = Pt(6)
            style.paragraph_format.line_spacing = 1.10

    for name, size, color, before, after in (
        ("Heading 1", 16, "2E74B5", 16, 8),
        ("Heading 2", 13, "2E74B5", 12, 6),
        ("Heading 3", 12, "1F4D78", 8, 4),
    ):
        style = doc.styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True


def add_cover(doc):
    p = doc.add_paragraph(style="normal")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("SHIKASTUDIO")
    style_run(r, size=12, bold=True, color="1F4D78")

    p = doc.add_paragraph(style="normal")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("GAME DESIGN DOCUMENT")
    style_run(r, size=18, bold=True, color="2E74B5")

    p = doc.add_paragraph(style="normal")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("KHU VƯỜN MÙA HÈ")
    style_run(r, size=28, bold=True, color="0B2545")

    add_para(doc)
    add_table(
        doc,
        ["THÔNG TIN", "GIÁ TRỊ"],
        [
            ["Product name", "KHU VƯỜN MÙA HÈ"],
            ["Project version", "Concept GDD"],
            ["Company", "ShikaStudio"],
            ["Engine", "Unity 2D hoặc engine 2D tương đương"],
            ["Source idea", "Core Gameplay Loop: khám phá khu vườn sau nhà"],
            ["Người thực hiện", "Shika"],
        ],
        [2800, 6560],
    )


def build_doc():
    doc = Document(TEMPLATE)
    clear_document_body(doc)
    set_document_defaults(doc)

    add_cover(doc)

    add_heading(doc, "MỤC LỤC", 1, page_break_before=True)
    add_para(doc, "1. Game Overview")
    add_para(doc, "2. Core Gameplay Mechanics")
    add_para(doc, "3. Narrative and World Context")
    add_para(doc, "4. Systems Design")
    add_para(doc, "5. Visual & Audio")
    add_para(doc, "6. Level Design")
    add_para(doc, "7. Monetization")
    add_para(doc, "8. Technical Requirements")
    add_para(doc, "9. Production Plan")

    add_heading(doc, "1. GAME OVERVIEW", 1, page_break_before=True)
    add_para(doc, "KHU VƯỜN MÙA HÈ là game khám phá phiêu lưu nhẹ 2D, nơi người chơi nhập vai một cậu bé bước ra khu vườn phía sau nhà vào một buổi chiều mùa hè. Camera dùng góc nhìn top-down nghiêng giống cảm giác di chuyển trong Among Us: thấy nhân vật từ trên xuống hơi lệch, thấy rõ mặt đất, vật cản và lối đi, nhưng vẫn giữ được chiều sâu bằng layer foreground/background. Không gian quen thuộc dần biến thành một thế giới rộng lớn dưới góc nhìn trẻ con: bụi cỏ giống khu rừng nhỏ, tảng đá như cánh cửa bí mật, gốc cây như một pháo đài và ánh đom đóm như tín hiệu từ một miền lạ.")
    add_para(doc, "Trọng tâm của game không nằm ở chiến đấu hay điểm số, mà ở chuỗi tò mò - tương tác - quan sát - bất ngờ. Người chơi phát hiện sinh vật, học hành vi của chúng qua phản ứng tự nhiên, mở ra lối đi mới và thu thập những mảnh ký ức tuổi thơ thông qua các vật nhỏ trong vườn.")

    add_heading(doc, "1.1 Title and high-level concept", 2)
    add_para(doc, "“Một khu vườn quen thuộc trở thành thế giới kỳ diệu”: player điều khiển nhân vật 2D di chuyển mượt trong một bản đồ nghiêng/top-down, đi qua các lối nhỏ, lật đá, vạch cỏ, trèo cây, đào đất mềm hoặc soi đèn vào vùng tối để tìm những bí mật nhỏ nhưng đáng nhớ.")

    add_heading(doc, "1.2 Core vision pillars", 2)
    for text in [
        "Small interactions, big wonder: mỗi hành động nhỏ cần tạo cảm giác có phản hồi, có đời sống và có khả năng dẫn đến điều bất ngờ.",
        "Living nature: sinh vật không phải collectible đứng yên; chúng quan sát, trốn, cuộn lại, bay đi, tụ lại hoặc đổi hành vi theo thời tiết và thời điểm.",
        "Childlike perspective: game ưu tiên cảm giác của một đứa trẻ hơn tính mô phỏng chính xác; khu vườn có thể rộng, bí ẩn và hơi phóng đại.",
        "Show, don’t lecture: người chơi học bằng quan sát, âm thanh, chuyển động và hệ quả của tương tác thay vì tutorial dài.",
        "Quiet emotional reward: phần thưởng chính là khoảnh khắc khám phá, ký ức được gợi lại và cảm giác muốn xem còn gì phía sau bụi cây tiếp theo.",
    ]:
        add_para(doc, text)

    add_heading(doc, "1.3 Genre and benchmark titles", 2)
    add_para(doc, "Game thuộc nhóm 2D exploration adventure, cozy discovery và immersive nature vignette. Benchmark camera/movement gần với Among Us ở cảm giác nhân vật đi mượt trong bản đồ 2D top-down nghiêng, nhưng nhịp chơi chậm hơn, tập trung vào quan sát thiên nhiên và tương tác môi trường thay vì social deduction.")

    add_heading(doc, "1.4 Platforms", 2)
    add_para(doc, "Tập trung: PC và WebGL cho vertical slice. Điều khiển ưu tiên WASD/keyboard và gamepad; mobile có thể khả thi vì góc nhìn 2D phù hợp virtual joystick, nhưng cần pass riêng cho touch interact và kích thước vật nhỏ.")

    add_heading(doc, "1.5 Target audience", 2)
    add_para(doc, "Người chơi casual thích khám phá nhẹ, không áp lực thắng thua.")
    add_para(doc, "Người chơi yêu thiên nhiên, tuổi thơ, atmosphere và các tương tác môi trường nhỏ.")
    add_para(doc, "Người chơi trẻ tuổi hoặc gia đình nếu phần sinh vật nguy hiểm được thể hiện mềm, không gây sợ quá mức.")

    add_heading(doc, "1.6 Unique selling points", 2)
    add_para(doc, "USP của KHU VƯỜN MÙA HÈ là biến một không gian đời thường thành sân chơi khám phá giàu phản ứng. Một bụi cỏ rung, tiếng động dưới tảng đá hoặc vệt sáng ban đêm đều có thể là đầu mối cho một sinh vật, một lối đi, một ký ức hoặc một bí mật môi trường.")

    add_heading(doc, "2. CORE GAMEPLAY MECHANICS", 1)
    add_heading(doc, "2.1 Core loop", 2)
    for text in [
        "Player di chuyển tự do, mượt theo 8 hướng trong khu vườn 2D để tìm dấu hiệu lạ: âm thanh, chuyển động nhỏ, ánh sáng, vết đất mềm hoặc vật cũ bị che khuất.",
        "Player chọn hành động tương tác phù hợp như lật đá, vạch bụi cỏ, trèo cây, đào đất mềm, chui qua khe nhỏ hoặc soi đèn.",
        "Môi trường phản hồi bằng sinh vật xuất hiện, vật thể dịch chuyển, lối đi mở ra, ký ức được gợi lại hoặc trạng thái khu vực thay đổi.",
        "Player quan sát hành vi sinh vật để học quy luật: đến quá gần thì chim non hoảng, giun chui xuống đất, rắn trốn, đom đóm tụ theo ánh sáng.",
        "Khám phá mới mở thêm vùng, công cụ hoặc ký ức, làm khu vườn dần lớn hơn và bí ẩn hơn trong cảm nhận của người chơi.",
    ]:
        add_para(doc, text)

    add_heading(doc, "2.2 Player objectives", 2)
    add_para(doc, "Mục tiêu chính: khám phá các khu vực của khu vườn, tìm những hiện tượng kỳ lạ và mở dần bản đồ ký ức của cậu bé.")
    add_para(doc, "Mục tiêu phụ: ghi nhận sinh vật trong sổ quan sát, tìm đồ vật cũ, hoàn thành chuỗi tương tác môi trường và chứng kiến các khoảnh khắc hiếm theo thời điểm.")
    add_para(doc, "Mục tiêu cảm xúc: tạo cảm giác “mình vừa phát hiện một điều chỉ mình biết”, thay vì chỉ tối ưu điểm số.")

    add_heading(doc, "2.3 Player actions", 2)
    add_table(
        doc,
        ["Action", "Input đề xuất", "Design purpose"],
        [
            ["Move / look around", "WASD hoặc left stick", "Di chuyển mượt 8 hướng kiểu Among Us, tự do đọc không gian, tìm dấu hiệu lạ và chọn hướng khám phá."],
            ["Interact", "E / face button", "Lật đá, vạch cỏ, nhặt vật nhỏ hoặc chạm nhẹ vào sinh vật/môi trường."],
            ["Climb / crawl", "Giữ E tại điểm hợp lệ", "Mở đường tới góc nhìn mới: cành cây, khe hàng rào, gầm bụi rậm."],
            ["Dig", "Giữ E với xẻng nhỏ", "Tác động vào đất mềm để tìm giun, hộp cũ, đường hầm hoặc dấu vết."],
            ["Flashlight / mirror", "Right mouse / trigger", "Soi vùng tối, dẫn hướng đom đóm hoặc tạo phản ứng ánh sáng."],
            ["Observe journal", "Tab / menu", "Xem ghi chú sinh vật, ký ức, thời điểm xuất hiện và gợi ý tự nhiên."],
        ],
        [2200, 2200, 4960],
    )

    add_heading(doc, "2.4 Progression layers", 2)
    add_para(doc, "Tiến trình mở rộng theo ba lớp: không gian, công cụ và ký ức. Ban đầu player chỉ tiếp cận sân cỏ, gốc cây và hiên sau; sau đó mở khu bụi rậm, nhà cây cũ, góc hàng rào, bồn hoa, đường hầm dưới rễ và phiên bản ban đêm của khu vườn.")

    add_heading(doc, "2.5 Movement and camera feel", 2)
    add_para(doc, "Movement phải là free movement 2D, không đi theo ô. Nhân vật tăng/giảm tốc nhẹ để cảm giác mềm, nhưng vẫn phản hồi nhanh khi đổi hướng. Collider nên dùng capsule/circle nhỏ hơn sprite để player luồn qua lối hẹp tự nhiên mà không mắc cạnh.")
    add_para(doc, "Camera là orthographic top-down nghiêng, bám theo player với smoothing nhẹ. Màn hình nên cho thấy đủ không gian phía trước hướng di chuyển để người chơi đọc bụi cỏ, đá, sinh vật và lối nhỏ. Layer foreground như lá cây, cành thấp hoặc cỏ cao có thể che một phần nhân vật để tạo chiều sâu, nhưng phải fade hoặc cắt alpha khi che khuất quá lâu.")

    add_heading(doc, "3. NARRATIVE AND WORLD CONTEXT", 1)
    add_heading(doc, "3.1 Narrative elevator pitch", 2)
    add_para(doc, "Một cậu bé ở nhà trong buổi chiều mùa hè bước ra vườn sau để tìm nguồn gốc của một âm thanh lạ. Càng đi sâu vào những góc tưởng như quen thuộc, cậu càng phát hiện khu vườn chứa nhiều điều hơn trí nhớ của mình: sinh vật nhỏ, dấu vết cũ, món đồ từng đánh mất và những ký ức gia đình nằm lẫn trong đất, lá và ánh hoàng hôn.")

    add_heading(doc, "3.2 World building elements", 2)
    for text in [
        "Nhà sau và hiên gỗ là vùng an toàn, nơi người chơi bắt đầu và quay lại để nhìn khu vườn thay đổi theo ngày.",
        "Bụi cỏ, tảng đá, gốc cây, bồn hoa, ao nhỏ, hàng rào và nhà cây cũ là các micro-biome có sinh vật và quy luật riêng.",
        "Các món đồ cũ như viên bi, xe đồ chơi, ảnh rách, dây diều hoặc hộp thiếc mở khóa ký ức ngắn thay vì cutscene dài.",
        "Thời tiết và thời điểm làm thay đổi thế giới: sau mưa có giun và ốc; chạng vạng có chim về tổ; ban đêm có đom đóm và bóng cây khiến lối đi quen trở nên lạ.",
    ]:
        add_para(doc, text)

    add_heading(doc, "3.3 Tone and mood", 2)
    add_para(doc, "Tông game ấm, tò mò, hơi bí ẩn nhưng không kinh dị. Nguy hiểm như rắn hoặc bóng tối được xử lý bằng căng thẳng nhẹ và khoảng cách an toàn, để giữ cảm giác tuổi thơ: vừa sợ một chút, vừa muốn nhìn thêm một chút.")

    add_heading(doc, "4. SYSTEMS DESIGN", 1)
    add_heading(doc, "4.1 Exploration signal system", 2)
    add_para(doc, "Mỗi khu vực cần có tín hiệu dẫn tò mò: chuyển động nhỏ, âm thanh, ánh sáng, dấu chân, đất mềm hoặc vật thể lệch khỏi bối cảnh. Tín hiệu không nên biến thành waypoint UI; nó phải nằm trong môi trường và đủ rõ để người chơi muốn tiến lại gần.")

    add_heading(doc, "4.2 Environment interaction matrix", 2)
    add_table(
        doc,
        ["Element", "Possible interactions", "Possible outcomes"],
        [
            ["Bụi cỏ rung", "Vạch cỏ, đứng yên quan sát", "Bọ nhảy ra, đường mòn nhỏ lộ ra, chim non phản ứng nếu player đến quá nhanh."],
            ["Tảng đá", "Lật lên, soi đèn dưới khe", "Bọ cuộn tròn, giun trốn, rắn nhỏ trượt đi hoặc lộ đường hầm."],
            ["Gốc cây", "Đào đất mềm, chui qua rễ", "Mở lối ngắn, tìm hộp thiếc cũ, thấy đàn kiến di chuyển."],
            ["Cành cây thấp", "Trèo lên, quan sát tổ chim", "Mở góc nhìn bản đồ, thấy tín hiệu ở khu vực xa hơn."],
            ["Vùng tối", "Soi đèn, chờ mắt quen", "Đom đóm tụ lại, vật phát sáng xuất hiện, bóng cây hé lối đi."],
        ],
        [1900, 2800, 4660],
    )

    add_heading(doc, "4.3 Creature behavior rules", 2)
    add_table(
        doc,
        ["Creature", "Behavior", "Player learning"],
        [
            ["Bọ cuộn", "Cuộn tròn khi bị chạm, mở ra sau vài giây yên tĩnh.", "Dạy player giảm tốc và quan sát thay vì spam tương tác."],
            ["Giun đất", "Chui xuống khi có rung động mạnh, xuất hiện nhiều sau mưa.", "Dạy liên hệ giữa thời tiết, âm thanh bước chân và sinh vật."],
            ["Chim non", "Hoảng khi player đến gần tổ, bình tĩnh nếu player đứng ở khoảng cách vừa đủ.", "Dạy giới hạn tiếp cận và quan sát từ xa."],
            ["Rắn nhỏ", "Trốn rất nhanh dưới đá hoặc bụi rậm, tạo cảnh báo âm thanh nhẹ.", "Tạo cảm giác khu vườn có nguy hiểm nhưng không trừng phạt nặng."],
            ["Đom đóm", "Tụ lại nơi có ánh sáng mềm, tản ra nếu bị soi quá gắt.", "Dạy điều khiển ánh sáng như công cụ dẫn hướng."],
        ],
        [1700, 3800, 3860],
    )

    add_heading(doc, "4.4 Failure, reset and persistence", 2)
    add_para(doc, "Game nên hạn chế thất bại cứng. Nếu player làm sinh vật sợ hoặc bỏ lỡ khoảnh khắc, hệ thống chỉ chuyển trạng thái: sinh vật trốn, âm thanh im lại, hoặc cần chờ thời điểm khác. Reset dùng cho tình huống kẹt vật lý; tiến trình khám phá, ký ức và ghi chú sinh vật được lưu tự động.")

    add_heading(doc, "4.5 UI and menu flow", 2)
    add_para(doc, "UI tối giản: sổ quan sát, icon công cụ đang cầm, nhắc tương tác theo ngữ cảnh và bản đồ ký ức dạng phác thảo tay. Main menu nên vào game nhanh; không dùng quest log dày đặc vì mục tiêu là cảm giác tự phát hiện.")

    add_heading(doc, "4.6 Audio hooks", 2)
    add_para(doc, "Âm thanh là hệ thống dẫn đường cốt lõi: tiếng cỏ xào xạc, đá cạ đất, chim gọi, côn trùng, gió qua lá, mưa nhỏ và tiếng gỗ nhà cây. SFX cần có biến thể để tương tác lặp lại vẫn tự nhiên.")

    add_heading(doc, "5. VISUAL & AUDIO", 1)
    add_heading(doc, "5.1 Art style reference", 2)
    add_para(doc, "Định hướng visual nên là 2D stylized, mềm, ấm, có tỉ lệ hơi phóng đại theo góc nhìn trẻ con. Camera nhìn top-down nghiêng như Among Us: nhân vật là sprite 2D có animation đi mượt, môi trường là nền vẽ/tiles có chiều sâu bằng layer, sorting order và parallax rất nhẹ. Khu vườn không cần rộng thật, nhưng cần nhiều lớp: foreground lá/cỏ, midground vật tương tác và background nhà/hàng rào để player luôn cảm thấy mình ở trong một góc đời thường.")

    add_heading(doc, "5.2 Readability requirements", 2)
    for text in [
        "Tín hiệu tương tác phải nổi bật bằng animation/âm thanh nhẹ, không cần outline UI quá mạnh.",
        "Nhân vật phải đọc rõ ở kích thước nhỏ, có animation idle/walk theo hướng và bóng đổ mềm để bám với mặt đất.",
        "Sinh vật nhỏ cần silhouette rõ khi di chuyển và trạng thái riêng khi sợ, tò mò hoặc trốn.",
        "Vùng nguy hiểm nhẹ như rắn, cành cao hoặc bóng tối phải có cảnh báo diegetic trước khi player đến quá gần.",
        "Ban ngày, chiều mưa và ban đêm cần khác nhau về màu, âm thanh và loại sinh vật xuất hiện.",
        "Vật gợi ký ức cần có framing riêng: bụi được phủi, ánh sáng rơi lên vật, camera hạ thấp hoặc âm thanh nền lắng xuống.",
    ]:
        add_para(doc, text)

    add_heading(doc, "5.3 Audio direction", 2)
    add_para(doc, "BGM nên thưa, nhiều khoảng trống, ưu tiên texture mùa hè: piano/guitar nhẹ, pad ấm, tiếng ve xa và ambience vườn. Khi player phát hiện điều kỳ diệu, âm nhạc chỉ cần nhô lên rất nhỏ để không phá cảm giác riêng tư.")

    add_heading(doc, "6. LEVEL DESIGN", 1)
    add_heading(doc, "6.1 Garden zone design", 2)
    add_para(doc, "Level không chia thành màn rời rõ ràng, mà chia thành các zone trong cùng khu vườn. Mỗi zone có một motif tương tác và một nhóm sinh vật chính, giúp player nhớ bằng cảm giác nơi chốn thay vì danh sách nhiệm vụ.")

    add_heading(doc, "6.2 Proposed progression", 2)
    add_table(
        doc,
        ["Phase", "Garden focus", "Design note"],
        [
            ["Phase 1", "Hiên sau, sân cỏ, vài tảng đá", "Dạy di chuyển, interact, quan sát sinh vật phản ứng và sổ ghi chú."],
            ["Phase 2", "Bụi rậm, bồn hoa, gốc cây", "Mở đào đất, đường hầm nhỏ, sinh vật theo thời tiết sau mưa."],
            ["Phase 3", "Cành cây, tổ chim, nhà cây cũ", "Mở trèo cao, góc nhìn mới, ký ức về đồ chơi/tuổi thơ."],
            ["Phase 4", "Chạng vạng và ban đêm", "Mở đèn pin/gương, đom đóm, ánh sáng và cảm giác khu vườn biến đổi."],
        ],
        [1500, 2800, 5060],
    )

    add_heading(doc, "6.3 Difficulty curve principles", 2)
    for text in [
        "Mỗi công cụ mới cần một khoảnh khắc an toàn để thử trước khi dùng trong chuỗi tương tác dài.",
        "Không khóa tiến trình bằng câu đố tối nghĩa; nếu player bỏ lỡ, môi trường nên có nhiều tín hiệu thay thế.",
        "Sinh vật hiếm nên tạo động lực quay lại theo thời điểm/thời tiết, không nên trở thành checklist gây áp lực.",
        "Ký ức nên được phân bố như phần thưởng cảm xúc sau các phát hiện có chủ đích.",
        "Ban đêm tăng bí ẩn bằng ánh sáng và âm thanh, không biến game thành horror.",
    ]:
        add_para(doc, text)

    add_heading(doc, "7. MONETIZATION", 1)
    add_para(doc, "Với scope và tông cảm xúc hiện tại, hướng phù hợp nhất là premium ngắn trên PC/Web hoặc bản WebGL miễn phí để portfolio/showcase. Không nên dùng quảng cáo ngắt quãng vì nó phá atmosphere. Nếu mở rộng mobile, có thể cân nhắc bản paid nhỏ hoặc free demo + unlock full game.")

    add_heading(doc, "8. TECHNICAL REQUIREMENTS", 1)
    add_heading(doc, "8.1 Proposed architecture", 2)
    add_para(doc, "Nên tách hệ thống thành các module: Player2DMovement, CameraFollow2D, ExplorationSignal, InteractionTarget, CreatureBehavior, TimeWeatherState, MemoryUnlock và Journal. Môi trường nên dùng data-driven interaction definitions để designer thêm điểm tương tác mà không cần hard-code từng bụi cỏ hoặc tảng đá.")

    add_heading(doc, "8.2 Estimated PC specs", 2)
    add_para(doc, "Do game là 2D stylized phạm vi nhỏ, yêu cầu phần cứng dự kiến thấp: laptop phổ thông, 4 GB RAM trở lên và trình duyệt hỗ trợ WebGL 2 nếu build web. Target frame rate nên là 60FPS trên desktop để movement/camera giống Among Us đủ mượt.")

    add_heading(doc, "8.3 Known technical risks", 2)
    for text in [
        "AI sinh vật cần đủ sống động nhưng không quá phức tạp; ưu tiên state machine rõ, trigger nhỏ và animation tốt.",
        "Tương tác môi trường nhiều điểm nhỏ dễ thành content debt; cần editor tool để gắn tag, trạng thái và điều kiện xuất hiện.",
        "Góc nhìn top-down nghiêng có thể làm vật nhỏ bị che bởi foreground; cần hệ fade/sorting rõ cho cỏ, đá, cành cây và sinh vật.",
        "Hệ thời tiết/thời điểm phải ảnh hưởng đủ rõ nhưng không khiến player phải chờ quá lâu.",
        "Âm thanh định hướng cần mix kỹ để tín hiệu quan trọng không bị ambience che mất.",
    ]:
        add_para(doc, text)

    add_heading(doc, "9. PRODUCTION PLAN", 1)
    add_heading(doc, "9.1 Definition of done for vertical slice", 2)
    for text in [
        "Một khu vườn nhỏ có 4 zone liên thông: sân cỏ, bụi rậm, gốc cây và góc tối/chạng vạng.",
        "Ít nhất 8 điểm tương tác môi trường có phản hồi riêng và không trùng cảm giác.",
        "Ít nhất 5 sinh vật có hành vi quan sát được: bọ, giun, chim non, rắn nhỏ, đom đóm.",
        "Sổ quan sát lưu sinh vật, vật ký ức và mô tả ngắn sau khi player tự phát hiện.",
        "Một chuỗi ký ức hoàn chỉnh dẫn player từ buổi chiều sang khoảnh khắc ban đêm đầu tiên.",
        "Audio ambience và SFX đủ để người chơi tìm được ít nhất một bí mật bằng âm thanh.",
    ]:
        add_para(doc, text)

    add_heading(doc, "9.2 Roadmap", 2)
    add_table(
        doc,
        ["Phase", "Goal", "Deliverables"],
        [
            ["Pass 1", "Playable exploration slice", "2D smooth movement, orthographic follow camera, interact system, 1 zone, 3 interaction types, journal basic."],
            ["Pass 2", "Living garden", "Creature behaviors, environment signals, time/weather triggers, first memory items."],
            ["Pass 3", "Wonder moments", "Night variant, fireflies, hidden tunnel, treehouse reveal, polished audio cues."],
            ["Pass 4", "Release prep", "Content QA, accessibility/readability pass, WebGL build, trailer captures, store/itch assets."],
        ],
        [1500, 2500, 5360],
    )

    doc.save(OUTPUT)


if __name__ == "__main__":
    build_doc()
    print(OUTPUT)
